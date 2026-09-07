from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware

from bson import ObjectId
from datetime import datetime, timedelta

import re

from jose import jwt, JWTError

from passlib.context import CryptContext

from fastapi import Depends

from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials

from openpyxl.styles import Font, PatternFill
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from io import BytesIO

import json
import smtplib
import uuid
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

from fastapi.staticfiles import StaticFiles
import shutil

from fastapi import UploadFile, File
import os


from database import (
    case_studies,
    feedbacks,
    feedback_requests,
    contact_enquiries,
    notifications,
    users,
    password_resets,
    testimonials,
    image_library,
    jobs,
    db
)

from utils.image_library_utils import (
    list_all_images,
    get_image_by_id,
    delete_image,
    can_delete_image,
    increase_usage,
    decrease_usage,
    get_image_by_filename
)


from models import (
    CaseStudy,
    Feedback,
    ContactEnquiry,
    LoginRequest,
    ChangePasswordRequest,
    ResetAdminPasswordRequest,
    Testimonial,
    FeedbackRequestEmail,
    Job,
    
)


def extract_case_study_images(data):
    images = []

    if data.get("architecture_image") and data["architecture_image"].get("image"):
        images.append(data["architecture_image"]["image"])

    if data.get("workflow_image") and data["workflow_image"].get("image"):
        images.append(data["workflow_image"]["image"])

    if data.get("product_image") and data["product_image"].get("image"):
        images.append(data["product_image"]["image"])

    return images



import os

def update_image_usage(images, slug):

    print("===================================")
    print("Slug:", slug)
    print("Images:", images)

    for image_path in images:

        print("Original Path:", image_path)

        filename = os.path.basename(image_path)

        print("Filename:", filename)

        image = get_image_by_filename(filename)

        print("Mongo Image:", image)

        if image:

            print("Increasing usage...")

            increase_usage(
                str(image["_id"]),
                slug
            )

        else:
            print("Image NOT FOUND")



def remove_image_usage(images, slug):

    print("Removing usage:", slug)

    for image_path in images:

        filename = os.path.basename(image_path)

        image = get_image_by_filename(filename)

        print("Found:", image)

        if image:

            decrease_usage(
                str(image["_id"]),
                slug
            )
            
            
app = FastAPI()

IMAGE_UPLOAD_FOLDER = "uploads/case-studies"

os.makedirs(
    IMAGE_UPLOAD_FOLDER,
    exist_ok=True
)

UPLOAD_FOLDER = os.getenv("UPLOAD_FOLDER", "uploads")

os.makedirs(
    os.path.join(UPLOAD_FOLDER, "case-studies"),
    exist_ok=True
)

app.mount(
    "/uploads",
    StaticFiles(directory=UPLOAD_FOLDER),
    name="uploads"
)


pwd_context = CryptContext(
    schemes=["bcrypt"],
    deprecated="auto"
)

security = HTTPBearer()



from dotenv import load_dotenv

load_dotenv()


SECRET_KEY = os.getenv("SECRET_KEY")

ALGORITHM = "HS256"


SMTP_SERVER = os.getenv("BREVO_SMTP_SERVER")
SMTP_PORT = int(os.getenv("BREVO_SMTP_PORT", 587))

SMTP_USER = os.getenv("BREVO_SMTP_USER")
SMTP_PASSWORD = os.getenv("BREVO_SMTP_PASSWORD")

SENDER_EMAIL = os.getenv("SENDER_EMAIL")
RECEIVER_EMAIL = os.getenv("RECEIVER_EMAIL")

WEBSITE_URL = os.getenv("WEBSITE_URL")

print("SMTP_SERVER:", SMTP_SERVER)
print("SMTP_PORT:", SMTP_PORT)

print("SENDER_EMAIL:", SENDER_EMAIL)
print("RECEIVER_EMAIL:", RECEIVER_EMAIL)

def hash_password(password):

    return pwd_context.hash(password)


def verify_password(
    plain_password,
    hashed_password
):

    return pwd_context.verify(
        plain_password,
        hashed_password
    )


def validate_password(password):

    if len(password) < 12:
        return False

    if not re.search(r"[A-Z]", password):
        return False

    if not re.search(r"[a-z]", password):
        return False

    if not re.search(r"\d", password):
        return False

    if not re.search(
        r"[!@#$%^&*(),.?\":{}|<>]",
        password
    ):
        return False

    return True



def create_access_token(data):

    payload = data.copy()

    payload["exp"] = (
        datetime.utcnow()
        + timedelta(hours=24)
    )

    return jwt.encode(
        payload,
        SECRET_KEY,
        algorithm=ALGORITHM
    )
    
    
@app.get("/create-default-users")
def create_default_users():

    users.delete_many({})

    users.insert_many([
        {
            "email": "developer@yuktic.com",
            "password": hash_password("AdminYuktic@123"),
            "role": "developer"
        },
        {
            "email": "admin@yuktic.com",
            "password": hash_password("AdminYuktic@123"),
            "role": "admin"
        }
    ])

    return {
        "success": True
    }
    
    
def get_current_user(
    credentials: HTTPAuthorizationCredentials = Depends(security)
):

    try:

        token = credentials.credentials

        payload = jwt.decode(
            token,
            SECRET_KEY,
            algorithms=[ALGORITHM]
        )

        return payload

    except JWTError:

        raise HTTPException(
            status_code=401,
            detail="Invalid or expired token"
        )
        
        
def developer_only(
    user=Depends(get_current_user)
):

    if user["role"] != "developer":

        raise HTTPException(
            status_code=403,
            detail="Developer access required"
        )

    return user
  
        
@app.post("/auth/login")
def login(data: LoginRequest):

    user = users.find_one({
        "email": data.email
    })

    if not user:

        raise HTTPException(
            status_code=401,
            detail="Invalid email or password"
        )

    if not verify_password(
        data.password,
        user["password"]
    ):
        raise HTTPException(
            status_code=401,
            detail="Invalid email or password"
        )

    token = create_access_token({
        "email": user["email"],
        "role": user["role"]
    })

    return {
        "success": True,
        "token": token,
        "role": user["role"],
        "email": user["email"]
    }
    
    
@app.post("/auth/change-password")
def change_password(
    data: ChangePasswordRequest,
    user=Depends(get_current_user)
):  
    
    db_user = users.find_one({
        "email": user["email"]
    })

    if not db_user:
        raise HTTPException(
            status_code=404,
            detail="User not found"
        )
        

    

    if not verify_password(
        data.currentPassword,
        db_user["password"]
    ):
        raise HTTPException(
            status_code=400,
            detail="Current password incorrect"
        )

    if not validate_password(
        data.newPassword
    ):
        raise HTTPException(
            status_code=400,
            detail="Password does not meet security requirements"
        )

    users.update_one(
        {
            "email": user["email"]
        },
        {
            "$set": {
                "password": hash_password(
                    data.newPassword
                )
            }
        }
    )
    

    return {
        "success": True,
        "message": "Password changed successfully"
    }
  
  
@app.post("/developer/reset-admin-password")
def reset_admin_password(
    data: ResetAdminPasswordRequest,
    user=Depends(developer_only)
):

    if not validate_password(
        data.newPassword
    ):
        raise HTTPException(
            status_code=400,
            detail="Password does not meet security requirements"
        )

    users.update_one(
        {
            "role": "admin"
        },
        {
            "$set": {
                "password": hash_password(
                    data.newPassword
                )
            }
        }
    )

    return {
        "success": True,
        "message": "Admin password reset successfully"
    }
    
    
@app.get("/auth/me")
def get_me(
    user=Depends(get_current_user)
):
    return user



def create_notification(
    title,
    message,
    notif_type,
    target_id,
    url
):
    notifications.insert_one({
        "title": title,
        "message": message,
        "type": notif_type,
        "targetId": str(target_id),
        "url": url,
        "createdAt": datetime.utcnow(),
        "isRead": False
    })
    

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.get("/")
def home():
    return {
        "status": "connected",
        "database": db.name
    }


@app.get("/test-db")
def test_db():
    return {
        "database": db.name,
        "collections": db.list_collection_names()
    }




def send_email(subject: str, html: str):

    try:

        message = MIMEMultipart("alternative")

        message["Subject"] = subject
        message["From"] = SENDER_EMAIL
        message["To"] = RECEIVER_EMAIL

        message.attach(
            MIMEText(html, "html")
        )

        with smtplib.SMTP(
            SMTP_SERVER,
            SMTP_PORT
        ) as server:

            server.starttls()

            server.login(
                SMTP_USER,
                SMTP_PASSWORD
            )

            server.sendmail(
                SENDER_EMAIL,
                RECEIVER_EMAIL,
                message.as_string()
            )

        print("✅ Email sent successfully")

    except Exception as e:

        print("❌ Email Error:", repr(e))

        raise
      
      

        
        
        
def send_feedback_request_email(
    to_email: str,
    subject: str,
    html: str
):

    try:

        message = MIMEMultipart("alternative")

        message["Subject"] = subject
        message["From"] = SENDER_EMAIL
        message["To"] = to_email

        message.attach(
            MIMEText(html, "html")
        )

        with smtplib.SMTP(
            SMTP_SERVER,
            SMTP_PORT
        ) as server:

            server.starttls()

            server.login(
                SMTP_USER,
                SMTP_PASSWORD
            )

            server.sendmail(
                SENDER_EMAIL,
                to_email,
                message.as_string()
            )

        print(f"✅ Feedback request sent to {to_email}")

    except Exception as e:

        print("❌ Feedback Request Email Error:", e)

        raise HTTPException(
            status_code=500,
            detail="Unable to send feedback request email."
        )






@app.post("/feedback/send-request")
def send_feedback_request(data: FeedbackRequestEmail):
    request_id = str(uuid.uuid4())
    html = f"""
    <!DOCTYPE html>
    <html lang="en">

    <head>

    <meta charset="UTF-8">
    <meta http-equiv="X-UA-Compatible" content="IE=edge">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">

    <title>Feedback Request</title>

    </head>

    <body style="
    margin:0;
    padding:0;
    background:#EEF3F1;
    font-family:Arial,Helvetica,sans-serif;
    -webkit-font-smoothing:antialiased;
    ">

    <table
    role="presentation"
    width="100%"
    border="0"
    cellpadding="0"
    cellspacing="0"
    style="
    width:100%;
    background:#EEF3F1;
    padding:32px 12px;
    ">

    <tr>

    <td align="center">

    <table
    role="presentation"
    width="760"
    border="0"
    cellpadding="0"
    cellspacing="0"
    style="
    width:760px;
    max-width:760px;
    background:#ffffff;
    border-radius:24px;
    overflow:hidden;
    box-shadow:0 8px 28px rgba(0,0,0,.08);
    ">

    <!-- ====================================================== -->
    <!-- HERO -->
    <!-- ====================================================== -->

   <tr>
        <td style="background-color:#EEF3F1; padding:40px 0 0 0;">
        
        <a
            href="{WEBSITE_URL}/feedback?requestId={request_id}"
            target="_blank"
            style="display:block; text-decoration:none; color:inherit;"
        >
            
            <table role="presentation" width="100%" border="0" cellpadding="0" cellspacing="0" style="background-color:#F5F9F7; border: 1px solid #DDE7E2; border-bottom: none; border-radius: 28px 28px 0 0; box-shadow: inset 0 2px 8px rgba(255,255,255,0.8); overflow:hidden;">
                <tr>
                    <td style="padding: 48px 32px;">
                        
                        <table role="presentation" width="100%" border="0" cellpadding="0" cellspacing="0">
                            <tr>
                                <td width="100%" align="center" valign="middle"
                                    style="padding:20px 50px; text-align:center;">
                                    
                                    <div style="margin:0 auto 32px; font-size:0; text-align:center;">
                                        <span style="color:#E84920; font-size:32px; font-weight:900; font-family:Arial, sans-serif; letter-spacing:-1px; vertical-align:middle;">YUKTIC</span>
                                        <span style="background-color:#E84920; color:#ffffff; font-size:16px; font-weight:bold; font-family:Arial, sans-serif; padding:4px 8px; border-radius:6px; vertical-align:middle; margin-left:6px; display:inline-block; line-height:1.2;">iT</span>
                                    </div>
                                    
                                   <h1 style="margin:0 0 6px; font-family:Arial,sans-serif; font-size:40px; line-height:1.1; font-weight:bold; color:#1C4931; text-align:center;">
                                        We Value
                                    </h1>
                                    <h1 style="margin:0; font-family:Arial,sans-serif; font-size:40px; line-height:1.1; font-weight:bold; color:#E84920; text-align:center;">
                                        Your Feedback
                                    </h1>
                                    
                                    <table role="presentation" border="0" cellpadding="0" cellspacing="0"
                                        style="margin:24px auto;">
                                        <tr>
                                            <td width="45" height="3" style="background-color:#1C4931; font-size:0; line-height:0;">&nbsp;</td>
                                        </tr>
                                    </table>
                                    
                                    <p style="margin:0 auto;
                                        max-width:520px;
                                        font-family:Arial,sans-serif;
                                        font-size:17px;
                                        line-height:1.8;
                                        color:#4A5550;
                                        text-align:center;">
                                        Your experience matters to us. We would greatly appreciate a few minutes of your time to share your thoughts about working with <strong style="color:#1C4931;">YUKTIC.</strong>
                                    </p>

                                </td>
                                
                                
                            </tr>
                        </table>

                    </td>
                </tr>
            </table>
            </a>
        </td>
    </tr>

    <!-- ====================================================== -->
    <!-- CONTENT WRAPPER -->
    <!-- ====================================================== -->

    <tr>

    <td
    style="
    padding:56px 56px 0 56px;
    background:#ffffff;
    ">

    <div
    style="
    font-size:17px;
    line-height:33px;
    color:#55625C;
    ">

    {data.message.replace(chr(10), "<br>")}

    </div>

    <div style="height:48px;"></div>

    <!-- ====================================================== -->
    <!-- SHARE EXPERIENCE CARD -->
    <!-- ====================================================== -->

    <table
    role="presentation"
    width="100%"
    border="0"
    cellpadding="0"
    cellspacing="0"
    style="
    background:#F7FBF9;
    border:1px solid #E5ECE8;
    border-radius:22px;
    overflow:hidden;
    ">

    <tr>

    <!-- LEFT -->

    <td
    width="175"
    align="center"
    valign="middle"
    style="
    padding:40px 28px;
    border-right:1px solid #DDE7E2;
    background:#F6FAF8;
    ">

    <table
    role="presentation"
    border="0"
    cellpadding="0"
    cellspacing="0">

    <tr>

    <td
    align="center"
    valign="middle"
    style="
    width:100px;
    height:100px;
    border-radius:50%;
    background:#EDF5F1;
    font-size:44px;
    line-height:100px;
    ">

    💬

    </td>

    </tr>

    </table>

    </td>

    <!-- RIGHT -->

    <td
    style="
    padding:38px 40px;
    ">

    <div
    style="
    font-size:24px;
    font-weight:700;
    color:#214D3D;
    line-height:34px;
    ">

    Share Your Experience

    </div>

    <div style="height:14px;"></div>

    <div
    style="
    font-size:16px;
    line-height:30px;
    color:#5E6A65;
    ">

    Your feedback will help us continue to innovate and deliver
    solutions that make a real difference.

    </div>

    <div style="height:28px;"></div>

    <table
    border="0"
    cellpadding="0"
    cellspacing="0">

    <tr>

    <td
    style="
    background:#F04A17;
    border-radius:8px;
    ">

    <a
    href="{WEBSITE_URL}/feedback?requestId={request_id}"
    style="
    display:inline-block;
    padding:17px 34px;
    font-size:17px;
    font-weight:bold;
    color:#ffffff;
    text-decoration:none;
    ">

    Share Your Feedback&nbsp;&nbsp;→

    </a>

    </td>

    </tr>

    </table>

    </td>

    </tr>

    </table>

    <div style="height:42px;"></div>
    <!-- ====================================================== -->
    <!-- THANK YOU INTRO -->
    <!-- ====================================================== -->

    <div
    style="
    font-size:17px;
    line-height:33px;
    color:#55625C;
    ">

    Every response we receive is carefully reviewed by our team.
    <br><br>

    Whether your experience was excellent or you believe there are
    areas where we can improve, we genuinely appreciate your honest
    feedback.
    <br><br>

    Our goal is to continuously enhance the quality of our services and
    provide the best possible experience for every client.

    </div>

    


    <div style="height:36px;"></div>

    <!-- ====================================================== -->
    <!-- FOOTER -->
    <!-- ====================================================== -->

    <table
    role="presentation"
    width="100%"
    border="0"
    cellpadding="0"
    cellspacing="0"
    style="
    background:#F8FAF9;
    border-radius:20px;
    ">

    <tr>

    <td
    align="center"
    style="
    padding:34px 26px;
    ">

    <table
    role="presentation"
    border="0"
    cellpadding="0"
    cellspacing="0">

    <tr>

    <td align="center">

    <a
    href="https://www.linkedin.com/company/yuktic/?viewAsMember=true"
    target="_blank"
    style="text-decoration:none;">

    <img
    src="https://cdn-icons-png.flaticon.com/512/3536/3536505.png"
    width="30"
    height="30"
    style="display:block;border:0;"
    alt="LinkedIn">

    </a>

    </td>

    <td width="18"></td>

    <td align="center">

    <a
    href="https://yuktic.com/"
    target="_blank"
    style="text-decoration:none;">

    <img
    src="https://cdn-icons-png.flaticon.com/512/1006/1006771.png"
    width="30"
    height="30"
    style="display:block;border:0;"
    alt="Website">

    </a>

    </td>

    <td width="18"></td>

    <td align="center">

    <a
    href="mailto:info@yuktic.com"
    style="text-decoration:none;">

    <img
    src="https://cdn-icons-png.flaticon.com/512/732/732200.png"
    width="30"
    height="30"
    style="display:block;border:0;"
    alt="Email">

    </a>

    </td>

    </tr>

    </table>

    <div style="height:28px;"></div>
    <table
    role="presentation"
    width="100%"
    border="0"
    cellpadding="0"
    cellspacing="0">

    <tr>

    <td
    style="
    border-top:1px solid #D8E2DC;
    font-size:1px;
    line-height:1px;
    ">

    &nbsp;

    </td>

    </tr>

    </table>

    <div style="height:26px;"></div>

    <div
    style="
    font-size:15px;
    font-weight:bold;
    line-height:24px;
    color:#2F4D43;
    ">

    © {datetime.now().year} YUKTIC. All rights reserved.

    </div>

    <div style="height:14px;"></div>

    <div
    style="
    font-size:14px;
    line-height:25px;
    color:#6C7773;
    max-width:540px;
    margin:0 auto;
    ">

    This is an automated email sent by YUKTIC.
    Please do not reply to this message unless you need assistance.

    </div>

    </td>

    </tr>

    </table>

    <!-- ====================================================== -->
    <!-- CLOSE CONTENT WRAPPER -->
    <!-- ====================================================== -->

    </td>

    </tr>

    <!-- ====================================================== -->
    <!-- END MAIN TABLE -->
    <!-- ====================================================== -->

    </table>

    </td>

    </tr>

    </table>

    </body>

    </html>
    """
    send_feedback_request_email(
            data.email,
            data.subject,
            html
        )
    
    feedback_requests.insert_one({
        "requestId": request_id,

        "email": data.email,

        "subject": data.subject,

        "message": data.message,

        "status": "Pending",

        "opened": False,

        "submitted": False,

        "openCount": 0,

        "openedAt": None,

        "submittedAt": None,

        "sentAt": datetime.utcnow()
    })

    return {
        "success": True,
        "message": "Feedback request sent successfully."
        
    }





@app.get("/feedback-requests")
def get_feedback_requests():

    requests = []

    for item in feedback_requests.find().sort("sentAt", -1):

        item["_id"] = str(item["_id"])

        requests.append(item)

    return requests


@app.delete("/feedback-requests/{id}")
def delete_feedback_request(id: str):

    feedback_requests.delete_one({
        "_id": ObjectId(id)
    })

    return {
        "success": True
    }
    
    
    
@app.put("/feedback-request/open/{request_id}")
def mark_feedback_open(request_id: str):

    result = feedback_requests.update_one(
        {
            "requestId": request_id
        },
        {
            "$set": {
                "opened": True,
                "status": "Opened",
                "openedAt": datetime.utcnow()
            },
            "$inc": {
                "openCount": 1
            }
        }
    )

    if result.matched_count == 0:
        raise HTTPException(
            status_code=404,
            detail="Request not found."
        )

    return {
        "success": True
    }


@app.put("/feedback-request/submit/{request_id}")
def mark_feedback_submitted(request_id: str):

    result = feedback_requests.update_one(
        {
            "requestId": request_id
        },
        {
            "$set": {
                "submitted": True,
                "status": "Submitted",
                "submittedAt": datetime.utcnow()
            }
        }
    )

    if result.matched_count == 0:
        raise HTTPException(
            status_code=404,
            detail="Request not found."
        )

    return {
        "success": True
    }
    
    

@app.delete("/feedback-requests")
def delete_all_feedback_requests():

    feedback_requests.delete_many({})

    return {
        "success": True
    }
    










    
    
    
def send_contact_email(data):

    name = data.get("name") or "-"
    email = data.get("email") or "-"
    phone = data.get("phone") or "-"
    subject = data.get("subject") or "Other"
    message = data.get("message") or "-"

    html = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="UTF-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">
        <title>New Contact Enquiry</title>
    </head>

    <body style="margin:0;padding:0;background:#eef3f0;font-family:Arial,Helvetica,sans-serif;">

        <table
            width="100%"
            cellpadding="0"
            cellspacing="0"
            style="background:#eef3f0;padding:40px 15px;"
        >
            <tr>
                <td align="center">

                    <table
                        width="720"
                        cellpadding="0"
                        cellspacing="0"
                        style="
                            width:720px;
                            max-width:720px;
                            background:#ffffff;
                            border-radius:18px;
                            overflow:hidden;
                        "
                    >

                        <!-- HEADER -->
                        <tr>
                            <td
                                style="
                                    background:linear-gradient(135deg,#2A6049,#234D3A);
                                    padding:45px 40px;
                                    text-align:center;
                                "
                            >

                                <div
                                    style="
                                        display:inline-block;
                                        background:white;
                                        padding:12px 26px;
                                        border-radius:50px;
                                        font-size:30px;
                                        font-weight:bold;
                                        color:#2A6049;
                                        letter-spacing:2px;
                                    "
                                >
                                    YUKTIC
                                </div>

                                <div style="height:24px;"></div>

                                <div
                                    style="
                                        font-size:34px;
                                        font-weight:bold;
                                        color:white;
                                    "
                                >
                                    New Contact Enquiry
                                </div>

                            </td>
                        </tr>


                        <!-- BODY -->
                        <tr>
                            <td style="padding:40px;background:white;">

                                <div
                                    style="
                                        font-size:27px;
                                        font-weight:bold;
                                        color:#2A6049;
                                    "
                                >
                                    Contact Information
                                </div>

                                <div style="height:25px;"></div>


                                <!-- NAME -->
                                <div
                                    style="
                                        background:#F7FAF8;
                                        border:1px solid #E3ECE7;
                                        border-radius:14px;
                                        padding:20px;
                                        margin-bottom:15px;
                                    "
                                >
                                    <div
                                        style="
                                            font-size:13px;
                                            font-weight:bold;
                                            color:#7B8C84;
                                            text-transform:uppercase;
                                            letter-spacing:1px;
                                        "
                                    >
                                        Name
                                    </div>

                                    <div
                                        style="
                                            margin-top:8px;
                                            font-size:18px;
                                            font-weight:bold;
                                            color:#234D3A;
                                        "
                                    >
                                        {name}
                                    </div>
                                </div>


                                <!-- EMAIL -->
                                <div
                                    style="
                                        background:#F7FAF8;
                                        border:1px solid #E3ECE7;
                                        border-radius:14px;
                                        padding:20px;
                                        margin-bottom:15px;
                                    "
                                >
                                    <div
                                        style="
                                            font-size:13px;
                                            font-weight:bold;
                                            color:#7B8C84;
                                            text-transform:uppercase;
                                            letter-spacing:1px;
                                        "
                                    >
                                        Email
                                    </div>

                                    <div
                                        style="
                                            margin-top:8px;
                                            font-size:18px;
                                            font-weight:bold;
                                            color:#234D3A;
                                        "
                                    >
                                        {email}
                                    </div>
                                </div>


                                <!-- PHONE -->
                                <div
                                    style="
                                        background:#F7FAF8;
                                        border:1px solid #E3ECE7;
                                        border-radius:14px;
                                        padding:20px;
                                        margin-bottom:15px;
                                    "
                                >
                                    <div
                                        style="
                                            font-size:13px;
                                            font-weight:bold;
                                            color:#7B8C84;
                                            text-transform:uppercase;
                                            letter-spacing:1px;
                                        "
                                    >
                                        Phone
                                    </div>

                                    <div
                                        style="
                                            margin-top:8px;
                                            font-size:18px;
                                            font-weight:bold;
                                            color:#234D3A;
                                        "
                                    >
                                        {phone}
                                    </div>
                                </div>


                                <!-- SUBJECT -->
                                <div
                                    style="
                                        background:#F7FAF8;
                                        border:1px solid #E3ECE7;
                                        border-radius:14px;
                                        padding:20px;
                                        margin-bottom:15px;
                                    "
                                >
                                    <div
                                        style="
                                            font-size:13px;
                                            font-weight:bold;
                                            color:#7B8C84;
                                            text-transform:uppercase;
                                            letter-spacing:1px;
                                        "
                                    >
                                        Subject
                                    </div>

                                    <div
                                        style="
                                            margin-top:8px;
                                            font-size:18px;
                                            font-weight:bold;
                                            color:#234D3A;
                                        "
                                    >
                                        {subject}
                                    </div>
                                </div>


                                <!-- MESSAGE -->
                                <div style="height:20px;"></div>

                                <div
                                    style="
                                        font-size:27px;
                                        font-weight:bold;
                                        color:#2A6049;
                                    "
                                >
                                    Message
                                </div>

                                <div style="height:15px;"></div>

                                <div
                                    style="
                                        background:#F8FBF9;
                                        border:1px solid #E4ECE8;
                                        border-left:6px solid #2A6049;
                                        border-radius:16px;
                                        padding:28px;
                                        font-size:16px;
                                        line-height:30px;
                                        color:#394843;
                                        white-space:pre-wrap;
                                        word-break:break-word;
                                    "
                                >
                                    {message}
                                </div>

                            </td>
                        </tr>


                        <!-- FOOTER -->
                        <tr>
                            <td
                                style="
                                    padding:30px;
                                    text-align:center;
                                    background:#ffffff;
                                "
                            >
                                <div
                                    style="
                                        font-size:13px;
                                        color:#8B9792;
                                    "
                                >
                                    © {datetime.now().year} YUKTIC
                                </div>

                                <div style="height:8px;"></div>

                                <div
                                    style="
                                        font-size:13px;
                                        color:#8B9792;
                                    "
                                >
                                    This is an automated notification email.
                                </div>
                            </td>
                        </tr>

                    </table>

                </td>
            </tr>
        </table>

    </body>
    </html>
    """

    send_email(
        "New Contact Enquiry | YUKTIC",
        html
    )

 



def send_feedback_email(data):
    # Pre-process data to handle missing fields and keep HTML clean
    ratings = data.get("ratings", {})
    
    # Client Info
    c_name = data.get("name") or "-"
    c_title = data.get("title") or "-"
    c_org = data.get("orgName") or "-"
    c_org2 = data.get("org2") or "-"
    c_country = data.get("country") or "-"
    c_org_type = data.get("orgType") or "-"
    c_linkedin = data.get("linkedin") or "-"
    
    # Project & Service Info
    p_name = data.get("projectName") or "-"
    p_duration = data.get("duration") or "-"
    types_str = ", ".join(data.get("types", [])) if data.get("types") else "-"
    services_str = ", ".join(data.get("services", [])) if data.get("services") else "-"
    
    # Executive Summary & Ratings
    r_overall = ratings.get("overall", "-")
    r_recommendation = data.get("recommendation") or "-"
    
    # Comments & Outcomes
    outcomes_str = ", ".join(data.get("outcomes", [])) if data.get("outcomes") else "-"
    c_impact = data.get("impact") or "-"
    c_worked_well = data.get("workedWell") or "-"
    c_improve = data.get("improve") or "-"
    
    # Consultant
    consultant_name = data.get("consultant") or "-"
    consultant_notes = data.get("consultantNotes") or "-"
    excellence_str = ", ".join(data.get("excellence", [])) if data.get("excellence") else "-"
    
    # Permissions & Consents
    publish_str = ", ".join(data.get("publish", [])) if data.get("publish") else "-"
    willing = data.get("willing") or "-"
    
    # Booleans to Yes/No
    def bool_to_yes_no(val):
        if val is True: return "Yes"
        if val is False: return "No"
        return "-"
        
    consent_process = bool_to_yes_no(data.get("consentProcess"))
    consent_removal = bool_to_yes_no(data.get("consentRemoval"))
    consent_contact = bool_to_yes_no(data.get("consentContact"))
    consent_publish = bool_to_yes_no(data.get("consentPublish"))
    
    record_id = data.get("_id", "Unknown")

    html = f"""<!DOCTYPE html>
    <html>
    <head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>YUKTIC Client Feedback</title>
    </head>
    <body style="margin:0; padding:0; background:#EEF3F1; font-family:Arial,Helvetica,sans-serif; color:#24332D;">
    <table width="100%" border="0" cellpadding="0" cellspacing="0" style="background:#EEF3F1; padding:30px 12px;">
    <tr>
    <td align="center">
    <table width="700" border="0" cellpadding="0" cellspacing="0" style="width:700px; max-width:700px; background:#FFFFFF; border-radius:18px; overflow:hidden; box-shadow:0 8px 25px rgba(0,0,0,.08);">
    <tr>
    <td style="background:#E64013; padding:34px 34px 28px; text-align:center;">
    <div style="display:inline-block; background:white; padding:10px 24px; border-radius:40px; font-size:26px; font-weight:bold; color:#E64013; letter-spacing:1px;">
    YUKTIC
    </div>
    <div style="height:18px;"></div>
    <div style="font-size:30px; font-weight:bold; color:white; line-height:38px;">
    Client Feedback Report
    </div>
    <div style="height:10px;"></div>
    <div style="font-size:15px; line-height:24px; color:#FFE8E0; max-width:500px; margin:auto;">
    A new client feedback form has been submitted through the YUKTIC website. Please review the feedback and follow up if required.
    </div>
    <div style="height:24px;"></div>
    <table align="center" border="0" cellpadding="0" cellspacing="0">
    <tr>
    <td style="background:rgba(255,255,255,.18); padding:10px 16px; border-radius:8px; font-size:13px; color:white;">
    📅 {datetime.now().strftime("%d %B %Y")}
    </td>
    <td width="10"></td>
    <td style="background:#2A6049; padding:10px 18px; border-radius:8px; font-size:13px; font-weight:bold; color:white;">
    NEW FEEDBACK
    </td>
    </tr>
    </table>
    </td>
    </tr>

    <tr>
    <td style="padding:28px;">
    <div style="font-size:24px; font-weight:bold; color:#E64013;">Client Information</div>
    <div style="height:18px;"></div>
    <table width="100%" border="0" cellpadding="0" cellspacing="0">
    <tr>
    <td width="50%" style="padding-right:8px; vertical-align:top;">
    <div style="background:#FFF7F4; border:1px solid #F6D8CF; border-radius:12px; padding:16px;">
    <div style="font-size:11px; font-weight:bold; letter-spacing:1px; text-transform:uppercase; color:#8B8B8B;">Client Name</div>
    <div style="margin-top:8px; font-size:18px; font-weight:bold; color:#E64013; line-height:26px; word-break:break-word; overflow-wrap:anywhere;">{c_name}</div>
    </div>
    </td>
    <td width="50%" style="padding-left:8px; vertical-align:top;">
    <div style="background:#F7FAF8; border:1px solid #E4ECE8; border-radius:12px; padding:16px;">
    <div style="font-size:11px; font-weight:bold; letter-spacing:1px; text-transform:uppercase; color:#8B8B8B;">Job Title</div>
    <div style="margin-top:8px; font-size:18px; font-weight:bold; color:#2A6049; line-height:26px; word-break:break-word; overflow-wrap:anywhere;">{c_title}</div>
    </div>
    </td>
    </tr>
    <tr>
    <td style="padding-top:14px; padding-right:8px; vertical-align:top;">
    <div style="background:#F7FAF8; border:1px solid #E4ECE8; border-radius:12px; padding:16px;">
    <div style="font-size:11px; font-weight:bold; letter-spacing:1px; text-transform:uppercase; color:#8B8B8B;">Organisation</div>
    <div style="margin-top:8px; font-size:17px; font-weight:bold; color:#2A6049; word-break:break-word; overflow-wrap:anywhere;">{c_org}</div>
    </div>
    </td>
    <td style="padding-top:14px; padding-left:8px; vertical-align:top;">
    <div style="background:#F7FAF8; border:1px solid #E4ECE8; border-radius:12px; padding:16px;">
    <div style="font-size:11px; font-weight:bold; letter-spacing:1px; text-transform:uppercase; color:#8B8B8B;">Organisation 2</div>
    <div style="margin-top:8px; font-size:17px; font-weight:bold; color:#2A6049; word-break:break-word; overflow-wrap:anywhere;">{c_org2}</div>
    </div>
    </td>
    </tr>
    <tr>
    <td style="padding-top:14px; padding-right:8px; vertical-align:top;">
    <div style="background:#F7FAF8; border:1px solid #E4ECE8; border-radius:12px; padding:16px;">
    <div style="font-size:11px; font-weight:bold; letter-spacing:1px; text-transform:uppercase; color:#8B8B8B;">Country</div>
    <div style="margin-top:8px; font-size:17px; font-weight:bold; color:#2A6049; word-break:break-word; overflow-wrap:anywhere;">{c_country}</div>
    </div>
    </td>
    <td style="padding-top:14px; padding-left:8px; vertical-align:top;">
    <div style="background:#F7FAF8; border:1px solid #E4ECE8; border-radius:12px; padding:16px;">
    <div style="font-size:11px; font-weight:bold; letter-spacing:1px; text-transform:uppercase; color:#8B8B8B;">Organisation Type</div>
    <div style="margin-top:8px; font-size:17px; font-weight:bold; color:#2A6049; word-break:break-word; overflow-wrap:anywhere;">{c_org_type}</div>
    </div>
    </td>
    </tr>
    <tr>
    <td colspan="2" style="padding-top:14px; vertical-align:top;">
    <div style="background:#F7FAF8; border:1px solid #E4ECE8; border-radius:12px; padding:16px;">
    <div style="font-size:11px; font-weight:bold; letter-spacing:1px; text-transform:uppercase; color:#8B8B8B;">LinkedIn Profile</div>
    <div style="margin-top:8px; font-size:15px; font-weight:bold; color:#2A6049; word-break:break-word; overflow-wrap:anywhere;">{c_linkedin}</div>
    </div>
    </td>
    </tr>
    </table>

    <div style="height:28px;"></div>

    <div style="font-size:24px; font-weight:bold; color:#E64013;">Project Information</div>
    <div style="height:18px;"></div>
    <table width="100%" border="0" cellpadding="0" cellspacing="0">
    <tr>
    <td width="50%" style="padding-right:8px; vertical-align:top;">
    <div style="background:#FFF7F4; border:1px solid #F6D8CF; border-radius:12px; padding:16px;">
    <div style="font-size:11px; font-weight:bold; letter-spacing:1px; text-transform:uppercase; color:#888;">Project Name</div>
    <div style="margin-top:8px; font-size:18px; font-weight:bold; color:#E64013; line-height:26px; word-break:break-word; overflow-wrap:anywhere;">{p_name}</div>
    </div>
    </td>
    <td width="50%" style="padding-left:8px; vertical-align:top;">
    <div style="background:#F7FAF8; border:1px solid #E4ECE8; border-radius:12px; padding:16px;">
    <div style="font-size:11px; font-weight:bold; letter-spacing:1px; text-transform:uppercase; color:#888;">Project Duration</div>
    <div style="margin-top:8px; font-size:18px; font-weight:bold; color:#2A6049; line-height:26px; word-break:break-word; overflow-wrap:anywhere;">{p_duration}</div>
    </div>
    </td>
    </tr>
    <tr>
    <td style="padding-top:14px; padding-right:8px; vertical-align:top;">
    <div style="background:#F7FAF8; border:1px solid #E4ECE8; border-radius:12px; padding:16px;">
    <div style="font-size:11px; font-weight:bold; letter-spacing:1px; text-transform:uppercase; color:#888;">Feedback Types</div>
    <div style="margin-top:8px; font-size:15px; font-weight:bold; color:#2A6049; line-height:22px; word-break:break-word; overflow-wrap:anywhere;">{types_str}</div>
    </div>
    </td>
    <td style="padding-top:14px; padding-left:8px; vertical-align:top;">
    <div style="background:#F7FAF8; border:1px solid #E4ECE8; border-radius:12px; padding:16px;">
    <div style="font-size:11px; font-weight:bold; letter-spacing:1px; text-transform:uppercase; color:#888;">Services Provided</div>
    <div style="margin-top:8px; font-size:15px; font-weight:bold; color:#2A6049; line-height:22px; word-break:break-word; overflow-wrap:anywhere;">{services_str}</div>
    </div>
    </td>
    </tr>
    </table>

    <div style="height:30px;"></div>

    <div style="font-size:24px; font-weight:bold; color:#E64013;">Executive Summary</div>
    <div style="height:18px;"></div>
    <table width="100%" border="0" cellpadding="0" cellspacing="0">
    <tr>
    <td width="50%" style="padding-right:6px; vertical-align:top;">
    <div style="background:#FFF7F4; border-top:4px solid #E64013; border-radius:12px; padding:18px; text-align:center;">
    <div style="font-size:11px; color:#888; font-weight:bold; letter-spacing:1px; text-transform:uppercase;">Overall Rating</div>
    <div style="margin-top:12px; font-size:32px; font-weight:bold; color:#E64013;">{r_overall}</div>
    <div style="margin-top:6px; font-size:13px; color:#666;">/ 5</div>
    </div>
    </td>
    <td width="50%" style="padding-left:6px; vertical-align:top;">
    <div style="background:#F7FAF8; border-top:4px solid #2A6049; border-radius:12px; padding:18px; text-align:center;">
    <div style="font-size:11px; color:#888; font-weight:bold; letter-spacing:1px; text-transform:uppercase;">Recommendation</div>
    <div style="margin-top:10px; font-size:15px; font-weight:bold; line-height:22px; color:#2A6049; word-break:break-word; overflow-wrap:anywhere;">{r_recommendation}</div>
    <div style="height:12px;"></div>
    </div>
    </td>
    </tr>
    </table>




    <div style="height:30px;"></div>

    <div style="font-size:24px; font-weight:bold; color:#E64013;">Performance Ratings</div>
    <div style="height:18px;"></div>
    <table width="100%" border="0" cellpadding="0" cellspacing="0" style="border-collapse:separate;border-spacing:10px; margin-left:-10px; margin-right:-10px;">
    <tr>
    <td width="50%" valign="top">
    <div style="background:#F8FAF9; border:1px solid #E4ECE8; border-radius:12px; padding:14px;">
    <table width="100%"><tr>
    <td style="font-size:14px; font-weight:bold; color:#444;">Professionalism</td>
    <td align="right" style="font-size:22px; font-weight:bold; color:#E64013;">{ratings.get("Professionalism", "-")}</td>
    </tr></table>
    </div>
    </td>
    <td width="50%" valign="top">
    <div style="background:#F8FAF9; border:1px solid #E4ECE8; border-radius:12px; padding:14px;">
    <table width="100%"><tr>
    <td style="font-size:14px; font-weight:bold; color:#444;">Communication</td>
    <td align="right" style="font-size:22px; font-weight:bold; color:#2A6049;">{ratings.get("Communication", "-")}</td>
    </tr></table>
    </div>
    </td>
    </tr>
    <tr>
    <td>
    <div style="background:#F8FAF9; border:1px solid #E4ECE8; border-radius:12px; padding:14px;">
    <table width="100%"><tr>
    <td style="font-size:14px; font-weight:bold; color:#444;">Technical Expertise</td>
    <td align="right" style="font-size:22px; font-weight:bold; color:#E64013;">{ratings.get("Technical expertise", "-")}</td>
    </tr></table>
    </div>
    </td>
    <td>
    <div style="background:#F8FAF9; border:1px solid #E4ECE8; border-radius:12px; padding:14px;">
    <table width="100%"><tr>
    <td style="font-size:14px; font-weight:bold; color:#444;">Leadership</td>
    <td align="right" style="font-size:22px; font-weight:bold; color:#2A6049;">{ratings.get("Leadership", "-")}</td>
    </tr></table>
    </div>
    </td>
    </tr>
    <tr>
    <td>
    <div style="background:#F8FAF9; border:1px solid #E4ECE8; border-radius:12px; padding:14px;">
    <table width="100%"><tr>
    <td style="font-size:14px; font-weight:bold; color:#444;">Problem Solving</td>
    <td align="right" style="font-size:22px; font-weight:bold; color:#E64013;">{ratings.get("Problem solving", "-")}</td>
    </tr></table>
    </div>
    </td>
    <td>
    <div style="background:#F8FAF9; border:1px solid #E4ECE8; border-radius:12px; padding:14px;">
    <table width="100%"><tr>
    <td style="font-size:14px; font-weight:bold; color:#444;">Delivery Quality</td>
    <td align="right" style="font-size:22px; font-weight:bold; color:#2A6049;">{ratings.get("Delivery quality", "-")}</td>
    </tr></table>
    </div>
    </td>
    </tr>
    <tr>
    <td>
    <div style="background:#F8FAF9; border:1px solid #E4ECE8; border-radius:12px; padding:14px;">
    <table width="100%"><tr>
    <td style="font-size:14px; font-weight:bold; color:#444;">Stakeholder Mgmt</td>
    <td align="right" style="font-size:22px; font-weight:bold; color:#E64013;">{ratings.get("Stakeholder management", "-")}</td>
    </tr></table>
    </div>
    </td>
    <td>
    <div style="background:#F8FAF9; border:1px solid #E4ECE8; border-radius:12px; padding:14px;">
    <table width="100%"><tr>
    <td style="font-size:14px; font-weight:bold; color:#444;">Value Delivered</td>
    <td align="right" style="font-size:22px; font-weight:bold; color:#2A6049;">{ratings.get("Value delivered", "-")}</td>
    </tr></table>
    </div>
    </td>
    </tr>
    <tr>
    <td>
    <div style="background:#F8FAF9; border:1px solid #E4ECE8; border-radius:12px; padding:14px;">
    <table width="100%"><tr>
    <td style="font-size:14px; font-weight:bold; color:#444;">Responsiveness</td>
    <td align="right" style="font-size:22px; font-weight:bold; color:#E64013;">{ratings.get("Responsiveness", "-")}</td>
    </tr></table>
    </div>
    </td>
    <td>
    <div style="background:#F8FAF9; border:1px solid #E4ECE8; border-radius:12px; padding:14px;">
    <table width="100%"><tr>
    <td style="font-size:14px; font-weight:bold; color:#444;">Knowledge Transfer</td>
    <td align="right" style="font-size:22px; font-weight:bold; color:#2A6049;">{ratings.get("Knowledge transfer", "-")}</td>
    </tr></table>
    </div>
    </td>
    </tr>
    </table>

    <div style="height:30px;"></div>

    <div style="font-size:24px; font-weight:bold; color:#E64013;">Client Comments</div>
    <div style="height:18px;"></div>

    <div style="background:#F8FBF9; border:1px solid #E4ECE8; border-left:5px solid #2A6049; border-radius:12px; padding:18px;">
    <div style="font-size:12px; font-weight:bold; text-transform:uppercase; letter-spacing:1px; color:#2A6049;">Positive Outcomes</div>
    <div style="margin-top:10px; font-size:15px; line-height:26px; color:#444; white-space:pre-wrap; word-break:break-word; overflow-wrap:anywhere;">{outcomes_str}</div>
    </div>
    <div style="height:14px;"></div>

    <div style="background:#FFF8F5; border:1px solid #F5D7CD; border-left:5px solid #E64013; border-radius:12px; padding:18px;">
    <div style="font-size:12px; font-weight:bold; text-transform:uppercase; letter-spacing:1px; color:#E64013;">Impact Delivered</div>
    <div style="margin-top:10px; font-size:15px; line-height:26px; color:#444; white-space:pre-wrap; word-break:break-word; overflow-wrap:anywhere;">{c_impact}</div>
    </div>
    <div style="height:14px;"></div>

    <div style="background:#F8FBF9; border:1px solid #E4ECE8; border-left:5px solid #2A6049; border-radius:12px; padding:18px;">
    <div style="font-size:12px; font-weight:bold; text-transform:uppercase; letter-spacing:1px; color:#2A6049;">What Worked Well</div>
    <div style="margin-top:10px; font-size:15px; line-height:26px; color:#444; white-space:pre-wrap; word-break:break-word; overflow-wrap:anywhere;">{c_worked_well}</div>
    </div>
    <div style="height:14px;"></div>

    <div style="background:#FFF8F5; border:1px solid #F5D7CD; border-left:5px solid #E64013; border-radius:12px; padding:18px;">
    <div style="font-size:12px; font-weight:bold; text-transform:uppercase; letter-spacing:1px; color:#E64013;">Areas For Improvement</div>
    <div style="margin-top:10px; font-size:15px; line-height:26px; color:#444; white-space:pre-wrap; word-break:break-word; overflow-wrap:anywhere;">{c_improve}</div>
    </div>

    <div style="height:30px;"></div>

    <div style="font-size:24px; font-weight:bold; color:#E64013;">Consultant Information</div>
    <div style="height:18px;"></div>
    <table width="100%" border="0" cellpadding="0" cellspacing="0">
    <tr>
    <td width="50%" style="padding-right:8px; vertical-align:top;">
    <div style="background:#F7FAF8; border:1px solid #E4ECE8; border-radius:12px; padding:16px;">
    <div style="font-size:11px; font-weight:bold; text-transform:uppercase; letter-spacing:1px; color:#888;">Consultant</div>
    <div style="margin-top:8px; font-size:17px; font-weight:bold; color:#2A6049; word-break:break-word; overflow-wrap:anywhere;">{consultant_name}</div>
    </div>
    </td>
    <td width="50%" style="padding-left:8px; vertical-align:top;">
    <div style="background:#FFF7F4; border:1px solid #F6D8CF; border-radius:12px; padding:16px;">
    <div style="font-size:11px; font-weight:bold; text-transform:uppercase; letter-spacing:1px; color:#888;">Areas Of Excellence</div>
    <div style="margin-top:8px; font-size:15px; line-height:24px; color:#444; word-break:break-word; overflow-wrap:anywhere;">{excellence_str}</div>
    </div>
    </td>
    </tr>
    <tr>
    <td colspan="2" style="padding-top:14px;">
    <div style="background:#F8FBF9; border-left:5px solid #2A6049; border-radius:12px; padding:18px;">
    <div style="font-size:12px; font-weight:bold; text-transform:uppercase; letter-spacing:1px; color:#2A6049;">Consultant Notes</div>
    <div style="margin-top:10px; font-size:15px; line-height:26px; color:#444; white-space:pre-wrap; word-break:break-word; overflow-wrap:anywhere;">{consultant_notes}</div>
    </div>
    </td>
    </tr>
    </table>

    <div style="height:28px;"></div>

    <div style="font-size:24px; font-weight:bold; color:#E64013;">Consent & Permissions</div>
    <div style="height:16px;"></div>
    <div style="background:#F7FAF8; border:1px solid #E4ECE8; border-radius:12px; padding:18px; font-size:15px; line-height:28px; color:#444; word-break:break-word; overflow-wrap:anywhere;">
    <b>Willing To Provide Testimonial:</b> {willing}<br><br>
    <b>Publication Level:</b> {publish_str}<br><br>
    <b>Data Processing Consent:</b> {consent_process}<br>
    <b>Data Removal Understanding:</b> {consent_removal}<br>
    <b>Contact Consent:</b> {consent_contact}<br>
    <b>Publishing Consent:</b> {consent_publish}
    </div>

    <div style="height:35px;"></div>
    <hr style="border:none; border-top:1px solid #E4ECE8;">
    <div style="height:30px;"></div>

    <div style="text-align:center;">
    <div style="font-size:26px; font-weight:bold; color:#E64013;">YUKTIC</div>
    <div style="height:12px;"></div>
    <div style="font-size:15px; line-height:26px; color:#6E7772; max-width:520px; margin:auto;">
    This email was automatically generated from the YUKTIC feedback portal. Please review the submitted feedback and follow up if any action is required.
    </div>
    <div style="height:22px;"></div>
    <div style="display:inline-block; background:#E64013; padding:12px 24px; border-radius:8px; color:white; font-size:14px; font-weight:bold;">
    Client Feedback Report
    </div>
    <div style="height:26px;"></div>
    <div style="font-size:12px; color:#8A9590; line-height:22px;">
    © {datetime.now().year} YUKTIC<br>
    Healthcare • Digital Transformation • Consultancy<br><br>
    Automated Notification | Record ID: {record_id}
    </div>
    </div>
    </td>
    </tr>
    </table>
    </td>
    </tr>
    </table>
    </body>
    </html>
    """

    send_email(
        "⭐ Client Feedback Report | YUKTIC",
        html
    )


@app.post("/case-studies")
def create_case_study(data: CaseStudy, user=Depends(get_current_user)):
    
    
    existing = case_studies.find_one({
        "slug": data.slug
    })

    if existing:
        raise HTTPException(
            status_code=409,
            detail="Slug already exists"
        )
        
        
    result = case_studies.insert_one(
        data.model_dump()
    )
    
    images = extract_case_study_images(
        data.model_dump()
    )

    update_image_usage(
        images,
        data.slug
    )
    
    create_notification(
        title="Case Study Published",
        message=data.title,
        notif_type="case-study",
        target_id=result.inserted_id,
        url="/list-edit-case-study"
    )
    

    return {
        "success": True,
        "inserted_id": str(result.inserted_id),
        "message": "Case Study Published Successfully"
    }


@app.post("/image-library/upload")
async def upload_image(
    file: UploadFile = File(...),
    user=Depends(get_current_user)
):

    extension = os.path.splitext(file.filename)[1]

    filename = f"{uuid.uuid4()}{extension}"

    filepath = os.path.join(
        IMAGE_UPLOAD_FOLDER,
        filename
    )

    with open(filepath, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    image_doc = {

        "filename": filename,

        "original_name": file.filename,

        "extension": extension,

        "size": os.path.getsize(filepath),

        "url": f"/uploads/case-studies/{filename}",

        "uploaded_at": datetime.utcnow(),

        "uploaded_by": user["email"],

        "usage_count": 0,

        "used_by": []

    }

    result = image_library.insert_one(image_doc)

    image_doc["_id"] = str(result.inserted_id)

    return image_doc

@app.get("/image-library")
def get_all_images(
    user=Depends(get_current_user)
):
    return list_all_images()


@app.get("/image-library/{image_id}")
def get_image(
    image_id: str,
    user=Depends(get_current_user)
):

    image = get_image_by_id(image_id)

    if not image:
        raise HTTPException(
            status_code=404,
            detail="Image not found"
        )

    image["_id"] = str(image["_id"])

    return image


@app.delete("/image-library/{image_id}")
def remove_image(
    image_id: str,
    user=Depends(get_current_user)
):

    if not can_delete_image(image_id):

        raise HTTPException(
            status_code=400,
            detail="Image is currently used by one or more case studies."
        )

    if not delete_image(image_id):

        raise HTTPException(
            status_code=404,
            detail="Image not found"
        )

    return {
        "success": True,
        "message": "Image deleted successfully."
    }


    

    
@app.get("/case-studies")
def get_case_studies():

    studies = []

    for item in case_studies.find(
        {"published": True},
        {"_id": 0}
    ):
        studies.append(item)

    return studies




@app.get("/debug")
def debug():
    return {
        "database": db.name,
        "collections": db.list_collection_names()
    }
    

@app.get("/admin/case-studies/{slug}")
def get_case_study_admin(slug: str, user=Depends(get_current_user)):

    study = case_studies.find_one(
        {
            "slug": slug
        },
        {
            "_id": 0
        }
    )

    if not study:
        raise HTTPException(
            status_code=404,
            detail="Case Study Not Found"
        )

    return study

@app.delete("/case-studies/{slug}")
def delete_case_study(
    slug: str,
    user=Depends(get_current_user)
):

    study = case_studies.find_one({
        "slug": slug
    })
    
    old_images = extract_case_study_images(study)

    remove_image_usage(
        old_images,
        study["slug"]
    )

    if not study:
        raise HTTPException(
            status_code=404,
            detail="Case Study Not Found"
        )

    old_images = extract_case_study_images(study)

    remove_image_usage(
        old_images,
        study["slug"]
    )

  
   

    result = case_studies.delete_one({
        "slug": slug
    })

    if result.deleted_count == 0:
        raise HTTPException(
            status_code=404,
            detail="Case Study Not Found"
        )

    create_notification(
        title="Case Study Deleted",
        message=slug,
        notif_type="delete",
        target_id=slug,
        url="/list-edit-case-study"
    )

    return {
        "success": True,
        "message": "Case Study Deleted Successfully"
    }
    
    
@app.put("/case-studies/{slug}")
def update_case_study(
    slug: str,
    data: CaseStudy,
    user=Depends(get_current_user)
):

    old_case = case_studies.find_one({
        "slug": slug
    })
    
    old_images = extract_case_study_images(old_case)

    new_images = extract_case_study_images(
        data.model_dump()
    )

    

    if not old_case:
        raise HTTPException(
            status_code=404,
            detail="Case Study Not Found"
        )
        
    existing = case_studies.find_one({
        "slug": data.slug
    })

    if existing and existing["slug"] != slug:
        raise HTTPException(
            status_code=409,
            detail="Slug already exists"
        )
   
    remove_image_usage(
        old_images,
        slug
    )

    update_image_usage(
        new_images,
        data.slug
    )
    

    result = case_studies.update_one(
        {
            "slug": slug
        },
        {
            "$set": data.model_dump()
        }
    )

    if result.matched_count == 0:
        raise HTTPException(
            status_code=404,
            detail="Case Study Not Found"
        )

    create_notification(
        title="Case Study Updated",
        message=slug,
        notif_type="update",
        target_id=slug,
        url="/list-edit-case-study"
    )

    return {
        "success": True,
        "message": "Case Study Updated Successfully"
    }
    
    
@app.get("/case-studies/{slug}")
def get_case_study(slug: str):

    study = case_studies.find_one(
        {
            "slug": slug,
            "published": True
        },
        {
            "_id": 0
        }
    )

    if not study:
        raise HTTPException(
            status_code=404,
            detail="Case Study Not Found"
        )

    return study

@app.get("/case-studies/check-slug/{slug}")
def check_slug(slug: str):

    study = case_studies.find_one(
        {"slug": slug},
        {"_id": 1}
    )

    return {
        "exists": study is not None
    }
    
    
@app.put("/case-studies/{slug}/feature")
def feature_case_study(slug: str, user=Depends(get_current_user)):

    # remove existing featured
    case_studies.update_many(
        {},
        {
            "$set": {
                "featured": False
            }
        }
    )

    result = case_studies.update_one(
        {
            "slug": slug
        },
        {
            "$set": {
                "featured": True
            }
        }
    )

    if result.matched_count == 0:
        raise HTTPException(
            status_code=404,
            detail="Case Study Not Found"
        )

    return {
        "success": True
    }
    


@app.get("/featured-case-study")
def get_featured_case_study():

    study = case_studies.find_one(
        {
            "featured": True,
            "published": True
        },
        {
            "_id": 0
        }
    )

    return study



@app.post("/feedback")
def create_feedback(data: dict):

    data["createdAt"] = datetime.utcnow()
    
    try:
        send_feedback_email(data)
    except Exception as e:
        print(e)

    result = feedbacks.insert_one(data)
    
    request_id = data.get("requestId")

    if request_id:

        feedback_requests.update_one(
            {
                "requestId": request_id
            },
            {
                "$set": {
                    "feedbackId": str(result.inserted_id),
                    "submitted": True,
                    "status": "Submitted",
                    "submittedAt": datetime.utcnow()
                }
            }
        )
        
    create_notification(
        title="New Feedback Received",
        message=f"{data.get('name', 'Anonymous')} submitted feedback",
        notif_type="feedback",
        target_id=result.inserted_id,
        url=f"/feedback-responses/{result.inserted_id}"
    )

    return {
        "success": True,
        "inserted_id": str(result.inserted_id),
        "message": "Feedback Submitted Successfully"
    }

@app.get("/feedback")
def get_feedbacks(user=Depends(get_current_user)):

    items = []

    for item in feedbacks.find():

        item["_id"] = str(item["_id"])

        items.append(item)

    return items

@app.get("/feedback/{feedback_id}")
def get_feedback(feedback_id: str, user=Depends(get_current_user)):



    feedback = feedbacks.find_one(
        {"_id": ObjectId(feedback_id)}
    )

    if not feedback:
        raise HTTPException(
            status_code=404,
            detail="Feedback Not Found"
        )

    feedback["_id"] = str(feedback["_id"])

    return feedback





@app.delete("/feedback/{feedback_id}")
def delete_feedback(feedback_id: str, user=Depends(get_current_user)):

    result = feedbacks.delete_one(
        {
            "_id": ObjectId(feedback_id)
        }
    )

    if result.deleted_count == 0:
        raise HTTPException(
            status_code=404,
            detail="Feedback Not Found"
        )
        
    create_notification(
        title="Feedback Deleted",
        message=feedback_id,
        notif_type="delete",
        target_id=feedback_id,
        url="/feedback-responses"
    )

    return {
        "success": True,
        "message": "Feedback Deleted Successfully"
    }
    
    
@app.delete("/feedback")
def delete_all_feedback(user=Depends(get_current_user)):

    result = feedbacks.delete_many({})

    return {
        "success": True,
        "deleted_count": result.deleted_count
    }
    

def testimonial_to_json(item):

    item["_id"] = str(item["_id"])

    return item

@app.post("/testimonials")
def create_testimonial(
    data: Testimonial,
    user=Depends(get_current_user)
):

    testimonial = data.model_dump()

    testimonial["createdAt"] = datetime.utcnow()

    result = testimonials.insert_one(testimonial)

    create_notification(
        title="Testimonial Created",
        message=testimonial["name"],
        notif_type="testimonial",
        target_id=result.inserted_id,
        url="/testimonials"
    )

    return {
        "success": True,
        "id": str(result.inserted_id)
    }
@app.get("/admin/testimonials")
def get_admin_testimonials(
    user=Depends(get_current_user)
):

    items = []

    for item in testimonials.find().sort(
        "createdAt",
        -1
    ):

        items.append(
            testimonial_to_json(item)
        )

    return items


@app.get("/admin/testimonials/{testimonial_id}")
def get_admin_testimonial(
    testimonial_id: str,
    user=Depends(get_current_user)
):

    item = testimonials.find_one(
        {
            "_id": ObjectId(testimonial_id)
        }
    )

    if not item:

        raise HTTPException(
            status_code=404,
            detail="Testimonial not found"
        )

    return testimonial_to_json(item)


@app.put("/admin/testimonials/{testimonial_id}")
def update_testimonial(
    testimonial_id: str,
    data: Testimonial,
    user=Depends(get_current_user)
):

    result = testimonials.update_one(
        {
            "_id": ObjectId(testimonial_id)
        },
        {
            "$set": data.model_dump()
        }
    )

    if result.matched_count == 0:

        raise HTTPException(
            status_code=404,
            detail="Testimonial not found"
        )

    create_notification(
        title="Testimonial Updated",
        message=data.name,
        notif_type="testimonial",
        target_id=testimonial_id,
        url="/testimonials"
    )

    return {
        "success": True
    }
    

@app.delete("/admin/testimonials/{testimonial_id}")
def delete_testimonial(
    testimonial_id: str,
    user=Depends(get_current_user)
):

    result = testimonials.delete_one(
        {
            "_id": ObjectId(testimonial_id)
        }
    )

    if result.deleted_count == 0:

        raise HTTPException(
            status_code=404,
            detail="Testimonial not found"
        )

    create_notification(
        title="Testimonial Deleted",
        message=testimonial_id,
        notif_type="delete",
        target_id=testimonial_id,
        url="/testimonials"
    )

    return {
        "success": True
    }
    
@app.get("/testimonials")
def get_public_testimonials():

    items = []

    for item in testimonials.find(
        {
            "published": True
        }
    ).sort(
        "createdAt",
        -1
    ):

        items.append(
            testimonial_to_json(item)
        )

    return items


@app.get("/export/feedback")
def export_feedback(user=Depends(get_current_user)):

    wb = Workbook()
    ws = wb.active
    ws.title = "Feedback"

    headers = [
        "Review Types",
        "Services",

        "Organisation Name",
        "Country",
        "Organisation Type",
        "Project Name",
        "Project Duration",

        "Overall Experience",
        "Professionalism",
        "Communication",
        "Technical Expertise",
        "Leadership",
        "Problem Solving",
        "Delivery Quality",
        "Stakeholder Management",
        "Value Delivered",
        "Responsiveness",
        "Knowledge Transfer",

        "NPS Score",
        "Work Again",

        "Positive Outcomes",
        "Impact Delivered",
        "What Worked Well",
        "Areas For Improvement",

        "Consultant Name",
        "Areas Of Excellence",
        "Consultant Notes",

        "Willing To Provide Testimonial",
        "Recommendation",
        "Publication Permissions",

        "Submitter Name",
        "Job Title",
        "Organisation",
        "LinkedIn",

        "Consent Process",
        "Consent Removal",
        "Consent Contact",
        "Consent Publish",

        "Created At"
    ]

    ws.append(headers)

    header_fill = PatternFill(
        start_color="2A6049",
        end_color="2A6049",
        fill_type="solid"
    )

    for cell in ws[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.fill = header_fill

    for item in feedbacks.find():

        ratings = item.get("ratings", {})

        ws.append([

            ", ".join(item.get("types", [])),
            ", ".join(item.get("services", [])),

            item.get("orgName", ""),
            item.get("country", ""),
            item.get("orgType", ""),
            item.get("projectName", ""),
            item.get("duration", ""),

            ratings.get("overall", ""),
            ratings.get("Professionalism", ""),
            ratings.get("Communication", ""),
            ratings.get("Technical expertise", ""),
            ratings.get("Leadership", ""),
            ratings.get("Problem solving", ""),
            ratings.get("Delivery quality", ""),
            ratings.get("Stakeholder management", ""),
            ratings.get("Value delivered", ""),
            ratings.get("Responsiveness", ""),
            ratings.get("Knowledge transfer", ""),

            item.get("nps", ""),
            item.get("again", ""),

            ", ".join(item.get("outcomes", [])),
            item.get("impact", ""),
            item.get("workedWell", ""),
            item.get("improve", ""),

            item.get("consultant", ""),
            ", ".join(item.get("excellence", [])),
            item.get("consultantNotes", ""),

            item.get("willing", ""),
            item.get("recommendation", ""),
            ", ".join(item.get("publish", [])),

            item.get("name", ""),
            item.get("title", ""),
            item.get("org2", ""),
            item.get("linkedin", ""),

            str(item.get("consentProcess", "")),
            str(item.get("consentRemoval", "")),
            str(item.get("consentContact", "")),
            str(item.get("consentPublish", "")),

            str(item.get("createdAt", ""))
        ])

    for column in ws.columns:

        max_length = 0
        column_letter = column[0].column_letter

        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass

        ws.column_dimensions[
            column_letter
        ].width = min(max_length + 5, 50)

    file = BytesIO()

    wb.save(file)

    file.seek(0)

    return StreamingResponse(
        file,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition":
            "attachment; filename=feedback.xlsx"
        }
    )

    

@app.post("/contact")
def create_contact(data: ContactEnquiry):

    enquiry = data.model_dump()

    enquiry["submittedAt"] = datetime.utcnow()
    enquiry["status"] = "new"

    send_contact_email(enquiry)

    result = contact_enquiries.insert_one(enquiry)

    create_notification(
        title="New Contact Request",
        message=f"{enquiry.get('name', 'Unknown')} submitted an enquiry",
        notif_type="contact",
        target_id=str(result.inserted_id),
        url="/contact"
    )

    return {
        "success": True,
        "message": "Enquiry submitted successfully",
        "id": str(result.inserted_id)
    }


@app.get("/admin/contact")
def get_contact_requests(
    user=Depends(get_current_user)
):

    items = []

    for item in contact_enquiries.find().sort("submittedAt", -1):

        item["_id"] = str(item["_id"])

        items.append(item)

    return items





@app.get("/admin/contact/{contact_id}")
def get_contact_request(
    contact_id: str,  
    user=Depends(get_current_user)
):

    item = contact_enquiries.find_one(
        {"_id": ObjectId(contact_id)}
    )

    if not item:
        raise HTTPException(
            status_code=404,
            detail="Contact request not found"
        )

    item["_id"] = str(item["_id"])

    return item



@app.delete("/admin/contact/{contact_id}")
def delete_contact_request(contact_id: str, user=Depends(get_current_user)):

    result = contact_enquiries.delete_one(
        {"_id": ObjectId(contact_id)}
    )

    if result.deleted_count == 0:
        raise HTTPException(
            status_code=404,
            detail="Contact request not found"
        )
        
    create_notification(
        title="Contact Request Deleted",
        message=contact_id,
        notif_type="delete",
        target_id=contact_id,
        url="/contact"
    )

    return {
        "success": True,
        "message": "Deleted successfully"
    }
    
    
    
    
    
    
@app.get("/export/contact")
def export_contact(user=Depends(get_current_user)):

    wb = Workbook()
    ws = wb.active
    ws.title = "Contacts"

    headers = [
        "First Name",
        "Last Name",
        "Job Title",
        "Organisation",

        "Email",
        "Phone",
        "Preferred Contact Method",

        "Organisation Type",
        "Organisation Size",

        "Service Required",
        "Project Stage",
        "Timescale",

        "Platform/System",
        "How Did You Hear About Us",

        "Project Overview",
        "Additional Information",

        "Submitted At",
        "Status"
    ]

    ws.append(headers)

    header_fill = PatternFill(
        start_color="2A6049",
        end_color="2A6049",
        fill_type="solid"
    )

    for cell in ws[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.fill = header_fill

    for item in contact_enquiries.find():

        ws.append([

            item.get("firstName", ""),
            item.get("lastName", ""),
            item.get("jobTitle", ""),
            item.get("organisation", ""),

            item.get("email", ""),
            item.get("phone", ""),
            item.get("contactMethod", ""),

            item.get("orgType", ""),
            item.get("orgSize", ""),

            item.get("service", ""),
            item.get("stage", ""),
            item.get("timescale", ""),

            item.get("platform", ""),
            item.get("source", ""),

            item.get("overview", ""),
            item.get("additional", ""),

            str(item.get("submittedAt", "")),
            item.get("status", "")
        ])

    for column in ws.columns:

        max_length = 0
        column_letter = column[0].column_letter

        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass

        ws.column_dimensions[
            column_letter
        ].width = min(max_length + 5, 60)

    file = BytesIO()

    wb.save(file)

    file.seek(0)

    return StreamingResponse(
        file,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition":
            "attachment; filename=contacts.xlsx"
        }
    )
    
    
@app.delete("/admin/contact")
def delete_all_contacts(
    user=Depends(get_current_user)
):

    result = contact_enquiries.delete_many({})

    create_notification(
        title="All Contact Requests Deleted",
        message="All enquiries removed",
        notif_type="delete",
        target_id="all",
        url="/contact"
    )

    return {
        "success": True,
        "deleted_count": result.deleted_count
    }
    
    
@app.get("/notifications")
def get_notifications(
    user=Depends(get_current_user)
):

    items = []

    for item in notifications.find().sort(
        "createdAt",
        -1
    ):

        item["_id"] = str(item["_id"])

        items.append(item)

    return items



@app.delete("/notifications/{notification_id}")
def delete_notification(notification_id: str,  user=Depends(get_current_user)):

    result = notifications.delete_one(
        {
            "_id": ObjectId(notification_id)
        }
    )

    return {
        "success": True
    }
    
    
    
@app.delete("/notifications")
def clear_notifications(user=Depends(get_current_user)):

    result = notifications.delete_many({})

    return {
        "success": True,
        "deleted_count": result.deleted_count
    }
    
    



@app.put("/testimonials/{id}/feature")
def feature_testimonial(id: str, user=Depends(get_current_user)):

    testimonials.update_many(
        {},
        {
            "$set": {
                "featured": False
            }
        }
    )

    result = testimonials.update_one(
        {
            "_id": ObjectId(id)
        },
        {
            "$set": {
                "featured": True
            }
        }
    )

    if result.matched_count == 0:
        raise HTTPException(
            status_code=404,
            detail="Testimonial Not Found"
        )

    return {
        "success": True
    }
    

@app.get("/testimonials/featured")
def get_featured_testimonial():

    testimonial = testimonials.find_one(
        {
            "featured": True
            
        }
    )

    if not testimonial:
        return None

    testimonial["_id"] = str(testimonial["_id"])

    return testimonial









# =========================================================
# JOBS
# =========================================================


@app.post("/jobs")
def create_job(
    data: Job,
    user=Depends(get_current_user)
):

    # Check duplicate slug
    existing = jobs.find_one({
        "slug": data.slug
    })

    if existing:
        raise HTTPException(
            status_code=409,
            detail="Slug already exists"
        )

    now = datetime.utcnow()

    job = data.model_dump()

    job["createdAt"] = now
    job["updatedAt"] = now
    job["createdBy"] = user["email"]
    job["updatedBy"] = user["email"]

    # Store exact publish date + time
    if job.get("published") is True:
        job["publishedAt"] = now
    else:
        job["publishedAt"] = None

    result = jobs.insert_one(job)

    create_notification(
        title="Job Created",
        message=data.title,
        notif_type="job",
        target_id=result.inserted_id,
        url="/jobs"
    )

    return {
        "success": True,
        "inserted_id": str(result.inserted_id),
        "message": "Job Created Successfully"
    }


# =========================================================
# ADMIN - GET ALL JOBS
# =========================================================

@app.get("/admin/jobs")
def get_admin_jobs(
    user=Depends(get_current_user)
):

    items = []

    for item in jobs.find().sort(
        [
            ("publishedAt", -1),
            ("createdAt", -1)
        ]
    ):

        item["_id"] = str(item["_id"])

        items.append(item)

    return items


# =========================================================
# ADMIN - GET SINGLE JOB
# =========================================================

@app.get("/admin/jobs/{slug}")
def get_admin_job(
    slug: str,
    user=Depends(get_current_user)
):

    job = jobs.find_one({
        "slug": slug
    })

    if not job:
        raise HTTPException(
            status_code=404,
            detail="Job Not Found"
        )

    job["_id"] = str(job["_id"])

    return job


# =========================================================
# UPDATE JOB
# =========================================================

@app.put("/jobs/{slug}")
def update_job(
    slug: str,
    data: Job,
    user=Depends(get_current_user)
):

    old_job = jobs.find_one({
        "slug": slug
    })

    if not old_job:
        raise HTTPException(
            status_code=404,
            detail="Job Not Found"
        )

    # Check if changing slug to an existing slug
    existing = jobs.find_one({
        "slug": data.slug
    })

    if existing and str(existing["_id"]) != str(old_job["_id"]):
        raise HTTPException(
            status_code=409,
            detail="Slug already exists"
        )

    now = datetime.utcnow()

    job = data.model_dump()

    job["updatedAt"] = now
    job["updatedBy"] = user["email"]

    # =====================================================
    # PUBLISH DATE / TIME LOGIC
    # =====================================================

    old_published = old_job.get(
        "published",
        False
    )

    new_published = job.get(
        "published",
        False
    )

    # Unpublished -> Published
    # Store NEW publish date/time
    if new_published and not old_published:

        job["publishedAt"] = now

    # Published -> Published
    # Keep ORIGINAL publish date/time
    elif new_published and old_published:

        job["publishedAt"] = old_job.get(
            "publishedAt"
        )

    # Published -> Unpublished
    # Remove publish date/time
    else:

        job["publishedAt"] = None

    result = jobs.update_one(
        {
            "_id": old_job["_id"]
        },
        {
            "$set": job
        }
    )

    if result.matched_count == 0:
        raise HTTPException(
            status_code=404,
            detail="Job Not Found"
        )

    create_notification(
        title="Job Updated",
        message=data.title,
        notif_type="update",
        target_id=str(old_job["_id"]),
        url="/jobs"
    )

    return {
        "success": True,
        "message": "Job Updated Successfully"
    }


# =========================================================
# DELETE JOB
# =========================================================

@app.delete("/jobs/{job_id}")
def delete_job(
    job_id: str,
    user=Depends(get_current_user)
):

    try:
        object_id = ObjectId(job_id)

    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Invalid Job ID"
        )

    job = jobs.find_one({
        "_id": object_id
    })

    if not job:
        raise HTTPException(
            status_code=404,
            detail="Job Not Found"
        )

    result = jobs.delete_one({
        "_id": object_id
    })

    if result.deleted_count == 0:
        raise HTTPException(
            status_code=404,
            detail="Job Not Found"
        )

    create_notification(
        title="Job Deleted",
        message=job.get(
            "title",
            "Job"
        ),
        notif_type="delete",
        target_id=job_id,
        url="/jobs"
    )

    return {
        "success": True,
        "message": "Job Deleted Successfully"
    }


# =========================================================
# CHECK JOB SLUG
# =========================================================

@app.get("/jobs/check-slug/{slug}")
def check_job_slug(
    slug: str,
    user=Depends(get_current_user)
):

    job = jobs.find_one(
        {
            "slug": slug
        },
        {
            "_id": 1
        }
    )

    return {
        "exists": job is not None
    }


# =========================================================
# PUBLISH / UNPUBLISH JOB
# =========================================================

@app.put("/jobs/{slug}/publish")
def publish_job(
    slug: str,
    user=Depends(get_current_user)
):

    job = jobs.find_one({
        "slug": slug
    })

    if not job:
        raise HTTPException(
            status_code=404,
            detail="Job Not Found"
        )

    now = datetime.utcnow()

    current_published = job.get(
        "published",
        False
    )

    # -----------------------------------------
    # UNPUBLISH
    # -----------------------------------------

    if current_published:

        jobs.update_one(
            {
                "_id": job["_id"]
            },
            {
                "$set": {
                    "published": False,
                    "publishedAt": None,
                    "updatedAt": now,
                    "updatedBy": user["email"]
                }
            }
        )

        create_notification(
            title="Job Unpublished",
            message=job.get(
                "title",
                "Job"
            ),
            notif_type="update",
            target_id=str(job["_id"]),
            url="/jobs"
        )

        return {
            "success": True,
            "published": False,
            "message": "Job Unpublished Successfully"
        }

    # -----------------------------------------
    # PUBLISH
    # -----------------------------------------

    jobs.update_one(
        {
            "_id": job["_id"]
        },
        {
            "$set": {
                "published": True,
                "publishedAt": now,
                "updatedAt": now,
                "updatedBy": user["email"]
            }
        }
    )

    create_notification(
        title="Job Published",
        message=job.get(
            "title",
            "Job"
        ),
        notif_type="job",
        target_id=str(job["_id"]),
        url="/jobs"
    )

    return {
        "success": True,
        "published": True,
        "publishedAt": now,
        "message": "Job Published Successfully"
    }


# =========================================================
# PUBLIC - GET ALL PUBLISHED JOBS
# LATEST PUBLISHED JOB FIRST
# =========================================================

@app.get("/jobs")
def get_jobs():

    items = []

    for item in jobs.find(
        {
            "published": True
        }
    ).sort(
        [
            ("publishedAt", -1),
            ("createdAt", -1)
        ]
    ):

        item["_id"] = str(item["_id"])

        items.append(item)

    return items


# =========================================================
# PUBLIC - GET SINGLE PUBLISHED JOB
# =========================================================

@app.get("/jobs/{slug}")
def get_public_job(
    slug: str
):

    job = jobs.find_one(
        {
            "slug": slug,
            "published": True
        }
    )

    if not job:
        raise HTTPException(
            status_code=404,
            detail="Job Not Found"
        )

    job["_id"] = str(job["_id"])

    return job